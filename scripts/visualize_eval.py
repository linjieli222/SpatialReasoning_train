#!/usr/bin/env python3
"""Visualize evaluation results as HTML with input images, questions, and predictions.

Usage:
    python visualize_eval.py --result_xlsx <path> --dataset <name> --output <path.html> [--num_samples 50] [--seed 42]

Examples:
    # TextCoT s2000 noEMA on td_path (50 random samples)
    python visualize_eval.py \
        --result_xlsx .../0002000_full_noema/eval/AI2ThorPT2P_td_path/bagel_mot/bagel_mot_AI2ThorPT2P_td_path_result.xlsx \
        --dataset AI2ThorPT2P_td_path \
        --output /gpfs/scrubbed/krishna/linjli/bagel_eval/textcot_s2000_noema_td_path_viz.html

    # VCoT with generated images
    python visualize_eval.py \
        --result_xlsx .../bagel_mot_vcot/bagel_mot_vcot_AI2ThorPT2P_dh_midpoint_result.xlsx \
        --dataset AI2ThorPT2P_dh_midpoint \
        --output /gpfs/scrubbed/krishna/linjli/bagel_eval/vcot_viz.html
"""

import argparse
import base64
import html
import os
import random
import re
import sys

import openpyxl


def load_dataset_images(dataset_name):
    """Load input images from the eval dataset as base64 strings.

    Uses the eval framework's build_dataset() to ensure correct subset/config mapping.
    """
    sys.path.insert(0, "/gpfs/home/linjli/source/SpatialReasoning_Eval")

    from vlmeval.dataset import build_dataset
    ds = build_dataset(dataset_name)
    if ds is None:
        print(f"Warning: Could not build dataset '{dataset_name}', skipping image loading")
        return {}

    images = {}
    for _, row in ds.data.iterrows():
        idx = int(row["index"])
        img_data = row.get("image", "")
        if img_data:
            # May have multiple images separated by <image_separator>
            images[idx] = img_data.split("<image_separator>") if "<image_separator>" in str(img_data) else [str(img_data)]
    return images


def load_results(result_xlsx):
    """Load result.xlsx into a list of dicts."""
    wb = openpyxl.load_workbook(result_xlsx, read_only=True)
    ws = wb.active
    headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        rows.append(dict(zip(headers, row)))
    wb.close()
    return rows


def extract_generated_images(prediction):
    """Extract [Image: /path/to/img.jpg] references from predictions."""
    if not prediction:
        return []
    return re.findall(r'\[Image:\s*([^\]]+)\]', str(prediction))


def image_to_base64(path):
    """Read an image file and return base64 data URI."""
    if not os.path.exists(path):
        return None
    with open(path, "rb") as f:
        data = f.read()
    ext = os.path.splitext(path)[1].lower()
    mime = {"jpg": "image/jpeg", "jpeg": "image/jpeg", "png": "image/png"}.get(ext.lstrip("."), "image/jpeg")
    return f"data:{mime};base64,{base64.b64encode(data).decode()}"


def format_prediction(prediction):
    """Format prediction text with highlighted tags."""
    if not prediction:
        return "<em>No prediction</em>"
    text = html.escape(str(prediction))
    # Highlight tags
    text = re.sub(r'&lt;think&gt;', '<span class="tag think">&lt;think&gt;</span>', text)
    text = re.sub(r'&lt;/think&gt;', '<span class="tag think">&lt;/think&gt;</span>', text)
    text = re.sub(r'&lt;answer&gt;', '<span class="tag answer">&lt;answer&gt;</span>', text)
    text = re.sub(r'&lt;/answer&gt;', '<span class="tag answer">&lt;/answer&gt;</span>', text)
    text = re.sub(r'&lt;image_start&gt;', '<span class="tag image">&lt;image_start&gt;</span>', text)
    text = re.sub(r'&lt;image_end&gt;', '<span class="tag image">&lt;image_end&gt;</span>', text)
    # Remove [Image: ...] references (we show them as actual images)
    text = re.sub(r'\[Image:\s*[^\]]+\]', '', text)
    # Preserve newlines
    text = text.replace('\n', '<br>')
    return text


def generate_html(samples, title, output_path):
    """Generate the HTML visualization file."""
    correct = sum(1 for s in samples if s.get("hit") == 1)
    total = len(samples)

    html_parts = [f"""<!DOCTYPE html>
<html>
<head>
<meta charset="utf-8">
<title>{html.escape(title)}</title>
<style>
    * {{ box-sizing: border-box; margin: 0; padding: 0; }}
    body {{ font-family: -apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, sans-serif; background: #f5f5f5; padding: 20px; }}
    .header {{ background: #1a1a2e; color: white; padding: 24px; border-radius: 12px; margin-bottom: 24px; }}
    .header h1 {{ font-size: 1.5em; margin-bottom: 8px; }}
    .header .stats {{ color: #aaa; font-size: 0.95em; }}
    .header .stats .acc {{ color: #4ecdc4; font-weight: bold; font-size: 1.1em; }}
    .sample {{ background: white; border-radius: 12px; margin-bottom: 20px; box-shadow: 0 2px 8px rgba(0,0,0,0.08); overflow: hidden; }}
    .sample-header {{ display: flex; align-items: center; justify-content: space-between; padding: 12px 20px; border-bottom: 1px solid #eee; }}
    .sample-header .idx {{ font-weight: 600; color: #555; }}
    .badge {{ padding: 4px 12px; border-radius: 20px; font-size: 0.85em; font-weight: 600; }}
    .badge.correct {{ background: #d4edda; color: #155724; }}
    .badge.wrong {{ background: #f8d7da; color: #721c24; }}
    .sample-body {{ display: flex; gap: 20px; padding: 20px; }}
    .images-col {{ flex: 0 0 auto; display: flex; flex-direction: column; gap: 10px; }}
    .images-col img {{ max-width: 350px; max-height: 350px; border-radius: 8px; border: 1px solid #ddd; }}
    .images-col .img-label {{ font-size: 0.8em; color: #888; text-align: center; }}
    .content-col {{ flex: 1; min-width: 0; }}
    .question {{ font-size: 1em; line-height: 1.5; margin-bottom: 12px; color: #333; }}
    .choices {{ display: grid; grid-template-columns: 1fr 1fr; gap: 6px; margin-bottom: 16px; }}
    .choice {{ padding: 8px 12px; border-radius: 6px; font-size: 0.9em; background: #f8f9fa; border: 1px solid #dee2e6; }}
    .choice.correct-answer {{ background: #d4edda; border-color: #28a745; font-weight: 600; }}
    .choice.predicted {{ border-color: #007bff; border-width: 2px; }}
    .choice.predicted.wrong-pred {{ border-color: #dc3545; background: #fff5f5; }}
    .prediction {{ background: #f8f9fa; border-radius: 8px; padding: 14px; font-size: 0.88em; line-height: 1.6; color: #444; max-height: 400px; overflow-y: auto; word-break: break-word; }}
    .prediction-label {{ font-size: 0.8em; color: #888; margin-bottom: 6px; font-weight: 600; text-transform: uppercase; }}
    .tag {{ font-weight: 700; padding: 1px 4px; border-radius: 3px; }}
    .tag.think {{ color: #6f42c1; background: #f3e8ff; }}
    .tag.answer {{ color: #0d6efd; background: #e7f1ff; }}
    .tag.image {{ color: #d63384; background: #fce4ec; }}
    .gen-images {{ display: flex; gap: 10px; margin-top: 12px; flex-wrap: wrap; }}
    .gen-images img {{ max-width: 300px; max-height: 300px; border-radius: 8px; border: 2px solid #d63384; }}
    .gen-images .img-label {{ font-size: 0.8em; color: #d63384; text-align: center; font-weight: 600; }}
    .filter-bar {{ display: flex; gap: 12px; margin-bottom: 16px; align-items: center; }}
    .filter-bar select, .filter-bar button {{ padding: 8px 14px; border-radius: 6px; border: 1px solid #ccc; font-size: 0.9em; }}
    .filter-bar button {{ background: #1a1a2e; color: white; cursor: pointer; border: none; }}
    .filter-bar button:hover {{ background: #16213e; }}
</style>
</head>
<body>
<div class="header">
    <h1>{html.escape(title)}</h1>
    <div class="stats">
        Showing {total} samples | Accuracy: <span class="acc">{correct}/{total} ({100*correct/total:.1f}%)</span>
    </div>
</div>
<div class="filter-bar">
    <select id="filterResult">
        <option value="all">All results</option>
        <option value="correct">Correct only</option>
        <option value="wrong">Wrong only</option>
    </select>
    <button onclick="applyFilter()">Filter</button>
</div>
<div id="samples">
"""]

    for s in samples:
        hit = s.get("hit", 0)
        idx = s.get("index", "?")
        question = html.escape(str(s.get("question", "")))
        answer = str(s.get("answer", ""))
        prediction_raw = str(s.get("prediction", ""))

        # Extract predicted answer letter
        pred_match = re.search(r'<answer>\s*([A-D])\s*</answer>', prediction_raw)
        pred_letter = pred_match.group(1) if pred_match else None

        badge_cls = "correct" if hit == 1 else "wrong"
        badge_text = "Correct" if hit == 1 else "Wrong"

        html_parts.append(f"""
<div class="sample" data-result="{badge_cls}">
    <div class="sample-header">
        <span class="idx">Sample #{idx}</span>
        <span class="badge {badge_cls}">{badge_text}</span>
    </div>
    <div class="sample-body">
        <div class="images-col">
""")

        # Input images
        for i, img_b64 in enumerate(s.get("input_images", [])):
            html_parts.append(f'            <div><img src="data:image/png;base64,{img_b64}" alt="Input {i+1}"><div class="img-label">Input Image {i+1}</div></div>\n')

        html_parts.append("        </div>\n        <div class=\"content-col\">\n")

        # Question
        html_parts.append(f'            <div class="question"><strong>Q:</strong> {question}</div>\n')

        # Answer choices
        html_parts.append('            <div class="choices">\n')
        for letter in ["A", "B", "C", "D"]:
            choice_text = html.escape(str(s.get(letter, "")))
            cls = "choice"
            if letter == answer:
                cls += " correct-answer"
            if letter == pred_letter:
                cls += " predicted"
                if letter != answer:
                    cls += " wrong-pred"
            html_parts.append(f'                <div class="{cls}"><strong>{letter}.</strong> {choice_text}</div>\n')
        html_parts.append('            </div>\n')

        # Prediction
        html_parts.append(f'            <div class="prediction-label">Model Prediction (GT: {answer}, Pred: {pred_letter or "?"})</div>\n')
        html_parts.append(f'            <div class="prediction">{format_prediction(prediction_raw)}</div>\n')

        # Generated images
        gen_images = s.get("generated_images", [])
        if gen_images:
            html_parts.append('            <div class="gen-images">\n')
            for i, img_data_uri in enumerate(gen_images):
                html_parts.append(f'                <div><img src="{img_data_uri}" alt="Generated {i+1}"><div class="img-label">Generated Image {i+1}</div></div>\n')
            html_parts.append('            </div>\n')

        html_parts.append("        </div>\n    </div>\n</div>\n")

    html_parts.append("""
</div>
<script>
function applyFilter() {
    const val = document.getElementById('filterResult').value;
    document.querySelectorAll('.sample').forEach(el => {
        if (val === 'all') el.style.display = '';
        else el.style.display = el.dataset.result === val ? '' : 'none';
    });
}
</script>
</body>
</html>
""")

    with open(output_path, "w") as f:
        f.write("".join(html_parts))

    size_mb = os.path.getsize(output_path) / 1024 / 1024
    print(f"Written: {output_path} ({size_mb:.1f} MB)")


def main():
    parser = argparse.ArgumentParser(description="Visualize eval results as HTML")
    parser.add_argument("--result_xlsx", required=True, help="Path to *_result.xlsx")
    parser.add_argument("--dataset", required=True, help="Dataset name (e.g. AI2ThorPT2P_td_path)")
    parser.add_argument("--output", required=True, help="Output HTML path")
    parser.add_argument("--num_samples", type=int, default=50, help="Number of samples to visualize")
    parser.add_argument("--seed", type=int, default=42, help="Random seed for sampling")
    parser.add_argument("--title", type=str, default=None, help="Title for the HTML page")
    args = parser.parse_args()

    print(f"Loading results from {args.result_xlsx}")
    results = load_results(args.result_xlsx)
    print(f"Loaded {len(results)} results")

    print(f"Loading dataset images for {args.dataset}")
    dataset_images = load_dataset_images(args.dataset)
    print(f"Loaded images for {len(dataset_images)} samples")

    # Subsample
    random.seed(args.seed)
    if args.num_samples > 0 and args.num_samples < len(results):
        results = random.sample(results, args.num_samples)
    results.sort(key=lambda x: x.get("index", 0))

    # Enrich samples with images
    samples = []
    for r in results:
        idx = int(r.get("index", 0))
        r["input_images"] = dataset_images.get(idx, [])

        # Extract generated images from prediction
        gen_paths = extract_generated_images(r.get("prediction"))
        gen_imgs = []
        for p in gen_paths:
            p = p.strip()
            data_uri = image_to_base64(p)
            if data_uri:
                gen_imgs.append(data_uri)
        r["generated_images"] = gen_imgs
        samples.append(r)

    title = args.title or f"Eval: {os.path.basename(args.result_xlsx).replace('_result.xlsx', '')}"
    print(f"Generating HTML for {len(samples)} samples...")
    generate_html(samples, title, args.output)


if __name__ == "__main__":
    main()
